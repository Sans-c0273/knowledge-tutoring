"""Programmatic fixture builders for the WU1 conversion/chunking tests.

Everything here is built in-process, offline, with libraries already in the
project's dependency set (python-pptx via markitdown[pptx], openpyxl, Pillow)
or with the standard library. PDF and DOCX are hand-assembled byte-for-byte so
no extra dev dependency (fpdf2 / python-docx) is required to run the suite —
DESIGN §1 lists those two as *optional* fixture-generation tools only.

Every builder returns the Path it wrote. All content is deterministic.

Import explicitly from tests:  ``from wu1_fixtures import make_pdf, ...``
(tests/ has no __init__.py; pytest's rootdir-relative import puts tests/ on
sys.path via the default "prepend" import mode).
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

THAI_PHRASE = "การเรียนรู้เชิงลึก"  # "deep learning"; must survive conversion (DESIGN §22)

# --------------------------------------------------------------------- text / md

MD_STRUCTURED = """Preamble line one, before any heading.
Preamble line two, still before any heading.

# Intro

Intro body: probability is the study of uncertainty.

## Terms

A sample space is the set of all possible outcomes.

## Scope

This unit covers discrete sample spaces only.

# Appendix

Appendix body: notation reference.
"""

MD_SETEXT = """Title Underlined
================

Body text under a setext heading.

Second Setext
-------------

More body text under the second setext heading.
"""

MD_THAI = f"""# บทนำ

{THAI_PHRASE} คือสาขาหนึ่งของการเรียนรู้ของเครื่อง

## ＡＩ Terms

Fullwidth letters above should be NFKC-normalised on read if the converter normalises.
"""


def make_text(path: Path, n_lines: int = 150, para_lines: int = 5) -> Path:
    """Plain text: `n_lines` lines grouped into paragraphs of `para_lines`, blank-line separated."""
    lines: list[str] = []
    i = 1
    while len(lines) < n_lines:
        for _ in range(para_lines):
            lines.append(f"line {i:03d} of the plain text fixture, about sample spaces.")
            i += 1
        lines.append("")
    path.write_text("\n".join(lines[:n_lines]) + "\n", encoding="utf-8")
    return path


def make_md(path: Path, text: str = MD_STRUCTURED) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def make_md_oversize_section(path: Path, n_paragraphs: int = 300, words_per_line: int = 12) -> Path:
    """One heading whose body is far larger than any sane unit (~6k+ words, 600 lines)."""
    body = []
    for p in range(n_paragraphs):
        body.append(" ".join(f"word{p:03d}_{w}" for w in range(words_per_line)))
        body.append(" ".join(f"tail{p:03d}_{w}" for w in range(words_per_line)))
        body.append("")
    path.write_text("# Big\n\n" + "\n".join(body), encoding="utf-8")
    return path


# ------------------------------------------------------------------------- PDF


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Minimal valid PDF (Helvetica, one text line per page). An empty string → empty page.

    Objects: 1 catalog, 2 pages, 3 font, then per page: page object + content stream.
    Cross-reference table offsets are computed exactly so strict readers accept it.
    """
    objs: list[bytes] = []
    n_pages = len(pages)
    # object numbers: 1=Catalog 2=Pages 3=Font, page i -> 4+2i (page), 5+2i (content)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        page_obj = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        ).encode()
        if text:
            stream = f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode("latin-1")
        else:
            stream = b""
        content_obj = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        objs.append(page_obj)
        objs.append(content_obj)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for num, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode())
    path.write_bytes(out.getvalue())
    return path


def make_corrupt_pdf(path: Path) -> Path:
    path.write_bytes(b"%PDF-1.4\nthis is not a pdf body at all\n%%EOF\n")
    return path


# ------------------------------------------------------------------------ PPTX


def make_pptx(path: Path, slides: list[tuple[str, str, str | None]]) -> Path:
    """slides = [(title, body, speaker_notes_or_None), ...] via python-pptx (markitdown[pptx] dep)."""
    from pptx import Presentation  # type: ignore
    from pptx.util import Inches  # type: ignore

    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content
    for title, body, notes in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        placeholder = slide.placeholders[1]
        placeholder.text = body
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


# ------------------------------------------------------------------------ DOCX

_CT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>"""

_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

_DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/></w:style>
</w:styles>"""

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _xml_escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _para(text: str, style: str | None) -> str:
    ppr = f"<w:pPr><w:pStyle w:val=\"{style}\"/></w:pPr>" if style else ""
    return f"<w:p>{ppr}<w:r><w:t xml:space=\"preserve\">{_xml_escape(text)}</w:t></w:r></w:p>"


def make_docx(path: Path, paragraphs: list[tuple[str | None, str]]) -> Path:
    """paragraphs = [(style, text)], style ∈ {None, "Heading1", "Heading2"}; real Word heading styles."""
    body = "".join(_para(text, style) for style, text in paragraphs)
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {_W}><w:body>{body}</w:body></w:document>'
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _RELS)
        z.writestr("word/_rels/document.xml.rels", _DOC_RELS)
        z.writestr("word/styles.xml", _STYLES)
        z.writestr("word/document.xml", document)
    return path


DOCX_WITH_HEADINGS: list[tuple[str | None, str]] = [
    ("Heading1", "1 Scope"),
    (None, "This specification covers the discrete case only."),
    ("Heading2", "1.2 Terms"),
    (None, "A sample space is the set of all possible outcomes of an experiment."),
    (None, "An event is a subset of the sample space."),
    ("Heading1", "2 Method"),
    (None, "Count favourable outcomes and divide by the total."),
]

DOCX_NO_HEADINGS: list[tuple[str | None, str]] = [(None, f"Paragraph {i:02d}: plain body text without any heading style.") for i in range(1, 41)]

DOCX_THAI: list[tuple[str | None, str]] = [
    ("Heading1", "บทนำ"),
    (None, f"{THAI_PHRASE} คือสาขาหนึ่งของการเรียนรู้ของเครื่อง"),
]


# ------------------------------------------------------------------------ XLSX


def make_xlsx(path: Path) -> Path:
    """Two sheets. Sheet1: header + 3 rows, two blank rows, then a second block. Sheet2: one block. One formula."""
    from openpyxl import Workbook  # type: ignore

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Term", "Definition", "Weight"])
    ws.append(["Sample space", "All possible outcomes", 3])
    ws.append(["Event", "Subset of the sample space", 2])
    ws.append(["Outcome", "Single element of the sample space", 1])
    ws.append([])
    ws.append([])
    ws.append(["Second block", "starts at row 7", None])
    ws.append(["Total", "=SUM(C2:C4)", None])
    ws2 = wb.create_sheet("Grades")
    ws2.append(["Student", "Score"])
    ws2.append(["A", 90])
    ws2.append(["B", 75])
    wb.save(str(path))
    return path


# ----------------------------------------------------------------------- image


def make_png(path: Path, size: tuple[int, int] = (2400, 1200)) -> Path:
    from PIL import Image, ImageDraw  # type: ignore

    w, h = size
    # Geometry is proportional to `size` so any (even tiny) canvas draws a valid
    # rectangle (x1 >= x0, y1 >= y0) — Pillow raises otherwise.
    margin_x, margin_y = max(1, w // 24), max(1, h // 12)
    stroke = max(1, min(w, h) // 150)
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    d.rectangle([margin_x, margin_y, max(margin_x, w // 2), max(margin_y, h - margin_y)], outline="black", width=stroke)
    d.line([w // 2, h // 2, max(w // 2, w - margin_x), h // 2], fill="black", width=stroke)
    img.save(str(path), format="PNG")
    return path


def make_jpeg(path: Path, size: tuple[int, int] = (800, 600)) -> Path:
    from PIL import Image  # type: ignore

    Image.new("RGB", size, (200, 30, 30)).save(str(path), format="JPEG")
    return path


# ------------------------------------------------------------------------- web

HTML_PAGE = """<!doctype html>
<html lang="en"><head><title>Probability Basics</title></head>
<body>
<nav><a href="/">Home</a> | <a href="/about">About</a></nav>
<article>
<h1>Probability Basics</h1>
<p>Probability quantifies uncertainty. This introductory article explains the vocabulary used throughout
the course, starting from the sample space and building up to events and their probabilities. Every
definition here is used later, so read it slowly and keep the examples in mind while you work.</p>
<h2>Sample Space</h2>
<p>The sample space is the set of all possible outcomes of an experiment. Rolling a fair die has a
sample space of six outcomes, one for each face. Tossing two coins has four outcomes, because the order
of the coins matters when we write them down as ordered pairs of heads and tails.</p>
<h2>Events</h2>
<p>An event is any subset of the sample space. The event "roll an even number" contains three outcomes.
Events can be combined with unions and intersections, and the probability of a union of disjoint events
is the sum of their probabilities. This additivity rule is the first axiom students meet.</p>
<h3>Complement</h3>
<p>The complement of an event is everything in the sample space that is not in the event. Its probability
is one minus the probability of the event, which is often the fastest route to an answer when the event
itself is awkward to count directly. <a href="https://example.org/more">Read more</a>.</p>
<img src="/diagram.png" alt="Venn diagram of two events">
</article>
<footer>Copyright notice and unrelated boilerplate that a good extractor discards.</footer>
</body></html>
"""

HTML_EMPTY = "<!doctype html><html><head><title>Empty</title></head><body><script>renderApp()</script></body></html>"

WEB_URL = "https://learn.example.edu/stats/probability-basics"


def make_html(path: Path, html: str = HTML_PAGE) -> Path:
    path.write_text(html, encoding="utf-8")
    return path


# ------------------------------------------------------- fake LLM boundary


class FakeCallStage:
    """Stand-in for `kg.llm.call_stage` (DESIGN §7.1) injected into the image converter.

    Mirrors the boundary's keyword surface — `stage, system, user_text, schema,
    images` (+ any extra kwargs such as cfg/ledger, accepted and ignored) — and
    returns a StageResult-shaped object whose `.data` is the next canned value.
    Records every call so tests can assert what crossed the boundary. Never
    touches a provider.
    """

    def __init__(self, canned: list[object] | None = None, *, raise_exc: BaseException | None = None) -> None:
        self.canned = list(canned or [])
        self.raise_exc = raise_exc
        self.calls: list[dict[str, object]] = []

    def __call__(self, stage: str, system: str, user_text: str, schema: type, images: list | None = None, **kw: object):
        self.calls.append(
            {"stage": stage, "system": system, "user_text": user_text, "schema": schema, "images": images, "extra": kw}
        )
        if self.raise_exc is not None:
            raise self.raise_exc
        if not self.canned:
            raise AssertionError("FakeCallStage exhausted: no canned response left")
        from types import SimpleNamespace

        return SimpleNamespace(
            data=self.canned.pop(0),
            usage=None,
            provider="fake",
            model_requested="fake-model",
            model_served=None,
            attempts=1,
        )


def canned_describe():
    """The canned `DescribeOutput` (DESIGN §4.3 / §8.0 shape) used by every image test."""
    from kg.schemas import DescribeOutput

    return DescribeOutput(
        kind="diagram",
        title="Venn diagram of two events",
        description="Two overlapping circles labelled A and B inside a rectangle representing the sample space.",
        elements=["Circle A", "Circle B", "Rectangle S"],
        text_visible=["A", "B", "S"],
        relationships=["A overlaps B", "A and B lie inside S"],
    )


def snapshot_tree(root: Path) -> dict[str, bytes]:
    """{relative posix path: content} for every file under `root` (for before/after diffs)."""
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob("*") if p.is_file()}
