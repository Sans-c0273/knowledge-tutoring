"""``.xlsx`` → Units via openpyxl (DESIGN §4.2, D3; R3 ``file.xlsx#Sheet1!A3:F40``).

Per sheet, rows are grouped into blocks separated by two or more empty rows.
Each block is a Markdown table whose header is the column letters and whose
first column is the Excel row number, so every cell in a unit is addressable.

Documented limitations: ``data_only=True`` reads cached formula values only — a
formula written by a program (no cache) is an empty cell; merged cells keep
their value in the top-left cell only; ``.xls`` is unsupported.
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from kg.convert._common import ConversionError, Unit, make_convert, nfkc, number_units, require_file, require_ooxml_part

EMPTY_ROWS_BETWEEN_BLOCKS = 2
_PLAIN_SHEET_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")


def sheet_ref(title: str) -> str:
    """Excel-style sheet reference: quoted when the name needs it (``'My Sheet'!A1``)."""
    return title if _PLAIN_SHEET_NAME.match(title) else "'" + title.replace("'", "''") + "'"


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    text = nfkc(str(value))
    return text.replace("\n", " ").replace("|", "\\|")


def _blocks(rows: list[tuple[object, ...]]) -> list[list[tuple[int, tuple[object, ...]]]]:
    """Group (row_number, values) into blocks split at >= EMPTY_ROWS_BETWEEN_BLOCKS empty rows."""
    blocks: list[list[tuple[int, tuple[object, ...]]]] = []
    current: list[tuple[int, tuple[object, ...]]] = []
    empties = 0
    for row_no, values in enumerate(rows, start=1):
        if all(_is_blank(v) for v in values):
            if not current:
                continue
            empties += 1
            current.append((row_no, values))
            if empties >= EMPTY_ROWS_BETWEEN_BLOCKS:
                blocks.append(current[:-empties])
                current, empties = [], 0
        else:
            empties = 0
            current.append((row_no, values))
    if current:
        blocks.append(current[: len(current) - empties] if empties else current)
    return [b for b in blocks if b]


def _table(block: list[tuple[int, tuple[object, ...]]], min_col: int, max_col: int) -> str:
    letters = [get_column_letter(c) for c in range(min_col, max_col + 1)]
    lines = ["| | " + " | ".join(letters) + " |", "|---|" + "---|" * len(letters)]
    for row_no, values in block:
        cells = [_cell_text(values[c - 1]) if c - 1 < len(values) else "" for c in range(min_col, max_col + 1)]
        lines.append(f"| {row_no} | " + " | ".join(cells) + " |")
    return "\n".join(lines)


def extract(src: Path) -> list[Unit]:
    src = require_ooxml_part(Path(src), "xl/workbook.xml", "Excel workbook")
    name = src.name
    try:
        wb = load_workbook(str(src), data_only=True, read_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError on damaged parts, ...
        raise ConversionError(f"{name}: cannot open workbook: {exc}") from exc
    units: list[Unit] = []
    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            for block in _blocks(rows):
                cols = [c + 1 for _, values in block for c, v in enumerate(values) if not _is_blank(v)]
                min_col, max_col = min(cols), max(cols)
                first_row, last_row = block[0][0], block[-1][0]
                a1 = f"{get_column_letter(min_col)}{first_row}:{get_column_letter(max_col)}{last_row}"
                locator = f"{name}#{sheet_ref(ws.title)}!{a1}"
                units.append(Unit("", locator, [ws.title], _table(block, min_col, max_col), "sheet_block", name))
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"{name}: cannot read workbook: {exc}") from exc
    finally:
        wb.close()
    if not units:
        raise ConversionError(f"{name}: workbook has no non-empty cells")
    return number_units(units)


convert = make_convert(extract)


def used_areas(src: Path) -> dict[str, tuple[int, int]]:
    """{sheet title: (max_row, max_column)} of the cells that hold a value."""
    wb = load_workbook(str(require_file(Path(src))), data_only=True, read_only=True)
    try:
        areas: dict[str, tuple[int, int]] = {}
        for ws in wb.worksheets:
            max_row = max_col = 0
            for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
                cols = [c + 1 for c, v in enumerate(values) if not _is_blank(v)]
                if cols:
                    max_row, max_col = row_no, max(max_col, max(cols))
            areas[ws.title] = (max_row, max_col)
        return areas
    finally:
        wb.close()


__all__ = ["EMPTY_ROWS_BETWEEN_BLOCKS", "convert", "extract", "sheet_ref", "used_areas"]
