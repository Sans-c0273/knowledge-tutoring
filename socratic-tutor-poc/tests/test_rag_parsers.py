"""Parser provenance tests (R3).

Fixtures are generated in-process with the same libraries the parsers use, so
the suite needs no binary files checked in and no network.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from socratic_tutor.domain.rag.parsers import (
    ParsedSegment,
    ParseError,
    check_fetch_target,
    detect_lang,
    parse_markdown,
    parse_path,
    parse_transcript_text,
    parse_url,
)

SEED_CORPUS = Path(__file__).resolve().parents[1] / "content" / "seed" / "rag-corpus"


def _minimal_pdf(pages: list[str], title: str) -> bytes:
    """Build a small text-bearing PDF by hand.

    No PDF writer in the dependency set can lay down text (pypdf edits existing
    pages, reportlab is not a dependency), so the fixture emits the objects and
    xref table directly. ASCII only — Helvetica/WinAnsi cannot encode Thai.
    """
    page_ids = [4 + 2 * index for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)

    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, text in enumerate(pages):
        page_id = page_ids[index]
        bodies.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        stream = f"BT /F1 12 Tf 40 720 Td ({text}) Tj ET".encode()
        bodies.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    info_id = len(bodies) + 1
    bodies.append(f"<< /Title ({title}) >>".encode())

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"

    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(bodies) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R /Info %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(bodies) + 1,
        info_id,
        xref_at,
    )
    return bytes(out)


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    path = tmp_path / "course-material.pdf"
    path.write_bytes(
        _minimal_pdf(
            [
                "Page 1: Inverse operations undo each other.",
                "Page 2: Balance both sides of the equation.",
            ],
            title="Algebra Basics",
        )
    )
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Two-Step Equations", level=1)
    document.add_paragraph("Undo the constant first, then the coefficient.")
    document.add_heading("Checking a Solution", level=1)
    document.add_paragraph("Substitute the value back into the original equation.")
    path = tmp_path / "worksheet.docx"
    document.save(str(path))
    return path


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    for title, body in [
        ("Distributive Property", "Multiply every term inside the parentheses."),
        ("Like Terms", "Same variable, same power."),
    ]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    path = tmp_path / "lesson-deck.pptx"
    presentation.save(str(path))
    return path


def test_detect_lang_uses_script_not_whitespace():
    assert detect_lang("Solve for x in 2x + 4 = 10") == "en"
    assert detect_lang("สมการสองขั้นตอน") == "th"
    # Thai explanations routinely embed Latin algebra and stay Thai.
    assert detect_lang("เช่น 2(x + 3) คือ 2 คูณ x") == "th"


def test_pdf_segments_carry_page_numbers(pdf_file: Path):
    segments = parse_path(pdf_file)

    assert len(segments) == 2
    assert [segment.page_or_slide_or_timestamp for segment in segments] == ["1", "2"]
    assert segments[0].source_ref == "Algebra Basics p.1"
    assert segments[1].source_ref == "Algebra Basics p.2"
    assert "Inverse operations" in segments[0].text


def test_docx_segments_carry_heading_provenance(docx_file: Path):
    segments = parse_path(docx_file)

    assert [segment.page_or_slide_or_timestamp for segment in segments] == [
        "Two-Step Equations",
        "Checking a Solution",
    ]
    assert segments[0].source_ref == "Worksheet — Two-Step Equations"
    assert "Undo the constant first" in segments[0].text


def test_pptx_segments_carry_slide_numbers(pptx_file: Path):
    segments = parse_path(pptx_file)

    assert len(segments) == 2
    assert [segment.page_or_slide_or_timestamp for segment in segments] == ["slide 1", "slide 2"]
    assert segments[0].source_ref == "Lesson Deck slide 1"
    assert "Distributive Property" in segments[0].text


def test_markdown_chunk_comments_supply_source_topic_and_lang():
    segments = parse_markdown(SEED_CORPUS / "ch2-solving-linear-equations.md")

    by_id = {segment.chunk_id: segment for segment in segments}
    assert by_id["c2-04"].source_ref == "Algebra Basics Ch.2 p.24"
    assert by_id["c2-04"].topic == "Two-Step Linear Equations"
    assert by_id["c2-04"].lang_hint == "th"
    assert by_id["c2-03"].lang_hint == "en"
    assert all(segment.source_ref for segment in segments)


def test_video_transcript_markdown_keeps_timestamp_refs():
    segments = parse_path(SEED_CORPUS / "video-transcript-distributive.md")

    refs = [segment.source_ref for segment in segments]
    assert "Distributive Property Video @ 02:30" in refs
    assert all("@" in ref for ref in refs)


def test_vtt_transcript_produces_timestamped_refs(tmp_path: Path):
    vtt = tmp_path / "distributive-property.vtt"
    vtt.write_text(
        "WEBVTT\n\n"
        "00:00:45.000 --> 00:00:49.000\n"
        "The distributive property multiplies every term inside.\n\n"
        "00:02:30.000 --> 00:02:34.000\n"
        "A common error is multiplying only the first term.\n",
        encoding="utf-8",
    )

    segments = parse_path(vtt)

    assert len(segments) == 1
    assert segments[0].source_ref == "Distributive Property @ 00:45"
    assert "first term" in segments[0].text


def test_srt_transcript_groups_cues_and_drops_cue_numbers(tmp_path: Path):
    srt = tmp_path / "lesson.srt"
    srt.write_text(
        "1\n00:01:05,000 --> 00:01:08,000\nCollect the variable terms first.\n\n"
        "2\n00:01:09,000 --> 00:01:12,000\nThen solve as a two-step equation.\n",
        encoding="utf-8",
    )

    segments = parse_transcript_text(srt.read_text(encoding="utf-8"), "Lesson")

    assert len(segments) == 1
    assert segments[0].source_ref == "Lesson @ 01:05"
    assert "1\n" not in segments[0].text


def test_unsupported_format_raises_parse_error(tmp_path: Path):
    unsupported = tmp_path / "diagram.png"
    unsupported.write_bytes(b"\x89PNG")

    with pytest.raises(ParseError, match="unsupported format"):
        parse_path(unsupported)


def test_missing_file_raises_parse_error(tmp_path: Path):
    with pytest.raises(ParseError, match="not a file"):
        parse_path(tmp_path / "nope.pdf")


def test_every_seed_segment_has_provenance():
    segments: list[ParsedSegment] = []
    for path in sorted(SEED_CORPUS.glob("*.md")):
        segments.extend(parse_path(path))

    assert segments
    assert all(segment.source_ref.strip() for segment in segments)


# ------------------------------------------------------------------------ SSRF


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1:8000/admin",
        "http://localhost/internal",
        "http://169.254.169.254/latest/meta-data/",
        "http://[::1]/",
        "http://0.0.0.0/",
    ],
)
def test_urls_pointing_at_this_machine_or_the_private_network_are_refused(url: str):
    """L13: whatever a URL fetches is indexed and later quoted to a student, so
    a fetch is content injection from the server's network position."""
    with pytest.raises(ParseError, match="refusing to fetch"):
        check_fetch_target(url)


@pytest.mark.parametrize("url", ["file:///etc/passwd", "gopher://example.com/", "ftp://host/x"])
def test_non_http_schemes_are_refused(url: str):
    with pytest.raises(ParseError, match="only http"):
        check_fetch_target(url)


def test_a_url_without_a_host_is_refused():
    with pytest.raises(ParseError, match="no host"):
        check_fetch_target("http:///nowhere")


_real_check = check_fetch_target


def test_a_redirect_into_the_private_network_is_refused(monkeypatch: pytest.MonkeyPatch):
    """A public URL that 302s to loopback must not slip past the first check."""
    import httpx

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "example.com":
            return httpx.Response(302, headers={"location": "http://127.0.0.1:8000/secret"})
        return httpx.Response(200, html="<html><body><p>internal</p></body></html>")

    monkeypatch.setattr(
        "socratic_tutor.domain.rag.parsers.check_fetch_target",
        lambda url: None if "example.com" in url else _real_check(url),
    )
    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)

    with pytest.raises(ParseError, match="refusing to fetch"):
        parse_url("http://example.com/lesson", client=client)
