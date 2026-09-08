"""Source parsers: file or URL → text segments carrying provenance.

Every segment leaves this module with a `source_ref` a student could follow back
to the original material ("Algebra Basics Ch.2 p.24", "Distributive Property
Video @ 02:30"). Tech Spec §4.2 requires a source reference on any factual
response, so a segment without provenance is never produced — the parsers fall
back to the document title rather than emitting an empty ref.

Video is transcript-only (Addendum §"Web UI scope" item 1): `.vtt`/`.srt`/`.md`/
`.txt` transcript files, timestamps parsed where present. No ASR, no decoding.
"""

from __future__ import annotations

import ipaddress
import re
import socket
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import httpx

#: Thai script block. Thai has no inter-word spaces, so language detection and
#: chunk-size estimation are codepoint-based, never whitespace-based.
_THAI_RANGE = re.compile(r"[฀-๿]")

#: `<!-- chunk: id | topic: X | source: Y | lang: z -->` — the seed corpus format.
#: Authored chunk boundaries are honoured as-is; see `chunker.chunk_segments`.
_CHUNK_COMMENT = re.compile(r"<!--\s*chunk:\s*(?P<body>.*?)\s*-->", re.DOTALL)

_MD_HEADING = re.compile(r"^(#{1,6})\s+(?P<title>.+?)\s*#*$", re.MULTILINE)

#: WebVTT / SRT cue timing line, e.g. `00:02:30.000 --> 00:02:34.500`.
_CUE_TIMING = re.compile(
    r"^(?P<start>(?:\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{1,3})\s*-->\s*"
    r"(?P<end>(?:\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{1,3})"
)

#: Inline `[00:02:30]` / `(2:30)` timestamps in plain-text transcripts.
_INLINE_TIMESTAMP = re.compile(r"^[\[(]?\s*(?P<ts>(?:\d{1,2}:)?\d{1,2}:\d{2})\s*[\])]?\s*[-–—]?\s*")

#: Transcript cues are grouped up to this many characters before a new segment
#: starts, so a 4-second caption does not become a 4-second retrieval chunk.
_CUE_GROUP_CHARS = 700

#: URL ingestion is restricted to the public web: anything fetched here is
#: indexed and later quoted to a student (see `check_fetch_target`).
ALLOWED_URL_SCHEMES = frozenset({"http", "https"})
MAX_REDIRECTS = 5
MAX_FETCH_BYTES = 20 * 1024 * 1024

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".text"}
TRANSCRIPT_SUFFIXES = {".vtt", ".srt"} | TEXT_SUFFIXES
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx"} | TRANSCRIPT_SUFFIXES


class ParseError(Exception):
    """A source could not be read or its format is not supported."""


@dataclass(frozen=True)
class ParsedSegment:
    """One provenance-bearing unit of source text.

    `page_or_slide_or_timestamp` is the raw locator ("24", "slide 3", "02:30");
    `source_ref` is the human-readable form that reaches the student. `topic` and
    `chunk_id` are populated only when the source declares them (seed-corpus
    markdown); everything else leaves them `None`.
    """

    text: str
    source_ref: str
    page_or_slide_or_timestamp: str | None = None
    lang_hint: str = "en"
    topic: str | None = None
    chunk_id: str | None = None


def detect_lang(text: str) -> str:
    """Return `"th"` if the text contains Thai script, else `"en"`.

    Presence, not majority: Thai explanations routinely embed Latin algebra
    ("2(x + 3) คือ 2 คูณ x") and are still Thai. A ratio test would split those
    sentences mid-thought.
    """
    return "th" if _THAI_RANGE.search(text) else "en"


def _clean(text: str) -> str:
    """Collapse trailing whitespace and runs of blank lines."""
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _title_from_path(path: Path) -> str:
    return re.sub(r"[-_]+", " ", path.stem).strip().title()


def _normalize_timestamp(raw: str) -> str:
    """`00:02:30.000` / `2:30` → `02:30`; keep hours only when non-zero."""
    stamp = raw.replace(",", ".").split(".")[0]
    parts = [int(p) for p in stamp.split(":")]
    if len(parts) == 3 and parts[0] == 0:
        parts = parts[1:]
    if len(parts) == 2:
        return f"{parts[0]:02d}:{parts[1]:02d}"
    return ":".join(f"{p:02d}" for p in parts)


def _segment(
    text: str,
    source_ref: str,
    locator: str | None = None,
    topic: str | None = None,
    chunk_id: str | None = None,
) -> ParsedSegment | None:
    """Build a segment, or `None` when the text is empty after cleaning."""
    body = _clean(text)
    if not body:
        return None
    return ParsedSegment(
        text=body,
        source_ref=source_ref,
        page_or_slide_or_timestamp=locator,
        lang_hint=detect_lang(body),
        topic=topic,
        chunk_id=chunk_id,
    )


# --------------------------------------------------------------------------- PDF


def parse_pdf(path: Path, title: str | None = None) -> list[ParsedSegment]:
    """One segment per page; `source_ref` = "<title> p.<n>"."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ParseError(f"could not read PDF {path.name}: {exc}") from exc

    doc_title = title or (reader.metadata or {}).get("/Title") or _title_from_path(path)
    doc_title = str(doc_title).strip() or _title_from_path(path)

    segments: list[ParsedSegment] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise ParseError(f"could not extract text from {path.name} p.{index}: {exc}") from exc
        seg = _segment(text, f"{doc_title} p.{index}", locator=str(index))
        if seg:
            segments.append(seg)
    return segments


# -------------------------------------------------------------------------- DOCX


def parse_docx(path: Path, title: str | None = None) -> list[ParsedSegment]:
    """One segment per heading section; `source_ref` = "<title> — <heading>".

    DOCX has no page model that survives parsing, so headings are the locator.
    Text before the first heading is attributed to the document itself.
    """
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError(f"could not read DOCX {path.name}: {exc}") from exc

    doc_title = title or (document.core_properties.title or "").strip() or _title_from_path(path)

    segments: list[ParsedSegment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        ref = f"{doc_title} — {heading}" if heading else doc_title
        seg = _segment("\n\n".join(buffer), ref, locator=heading)
        if seg:
            segments.append(seg)
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and (paragraph.style.name or "").startswith("Heading"):
            flush()
            heading = text
            continue
        buffer.append(text)
    flush()

    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        seg = _segment(
            "\n".join(rows),
            f"{doc_title} — table {table_index}",
            locator=f"table {table_index}",
        )
        if seg:
            segments.append(seg)
    return segments


# -------------------------------------------------------------------------- PPTX


def parse_pptx(path: Path, title: str | None = None) -> list[ParsedSegment]:
    """One segment per slide (shape text + speaker notes); `source_ref` = "<title> slide <n>"."""
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ParseError(f"could not read PPTX {path.name}: {exc}") from exc

    doc_title = (
        title or (presentation.core_properties.title or "").strip() or _title_from_path(path)
    )

    segments: list[ParsedSegment] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        seg = _segment(
            "\n\n".join(parts),
            f"{doc_title} slide {index}",
            locator=f"slide {index}",
        )
        if seg:
            segments.append(seg)
    return segments


# ---------------------------------------------------------------- Markdown / text


def _parse_chunk_comment(body: str) -> dict[str, str]:
    """`id | topic: X | source: Y | lang: z` → dict, first field being the id.

    Values may contain colons ("… Video @ 02:30"), so each field splits once.
    """
    fields = [part.strip() for part in body.split("|") if part.strip()]
    meta: dict[str, str] = {}
    if fields:
        meta["id"] = fields[0]
    for field in fields[1:]:
        if ":" not in field:
            continue
        key, value = field.split(":", 1)
        meta[key.strip().lower()] = value.strip()
    return meta


def parse_markdown_text(text: str, title: str) -> list[ParsedSegment]:
    """Parse markdown, honouring `<!-- chunk: ... -->` metadata when present.

    Declared chunks win over structural splitting: the comment carries the
    authored topic, source ref and language, and the chunker does not merge
    across segment boundaries, so an authored chunk survives intact.
    """
    text = text.replace("\r\n", "\n")
    matches = list(_CHUNK_COMMENT.finditer(text))
    if matches:
        segments: list[ParsedSegment] = []
        for position, match in enumerate(matches):
            meta = _parse_chunk_comment(match.group("body"))
            end = matches[position + 1].start() if position + 1 < len(matches) else len(text)
            body = text[match.end() : end]
            seg = _segment(
                body,
                meta.get("source") or title,
                locator=meta.get("source"),
                topic=meta.get("topic"),
                chunk_id=meta.get("id"),
            )
            if seg:
                declared = (meta.get("lang") or "").lower()
                if declared in {"th", "en"}:
                    seg = ParsedSegment(
                        text=seg.text,
                        source_ref=seg.source_ref,
                        page_or_slide_or_timestamp=seg.page_or_slide_or_timestamp,
                        lang_hint=declared,
                        topic=seg.topic,
                        chunk_id=seg.chunk_id,
                    )
                segments.append(seg)
        return segments

    # No authored chunks: split on headings, which are the natural boundaries.
    doc_title = title
    heading_matches = list(_MD_HEADING.finditer(text))
    if heading_matches and heading_matches[0].group(1) == "#":
        doc_title = heading_matches[0].group("title").strip()

    segments = []
    if not heading_matches:
        seg = _segment(text, doc_title)
        return [seg] if seg else []

    preamble = _segment(text[: heading_matches[0].start()], doc_title)
    if preamble:
        segments.append(preamble)

    for position, match in enumerate(heading_matches):
        heading = match.group("title").strip()
        end = (
            heading_matches[position + 1].start()
            if position + 1 < len(heading_matches)
            else len(text)
        )
        body = text[match.end() : end]
        ref = doc_title if heading == doc_title else f"{doc_title} — {heading}"
        seg = _segment(body, ref, locator=heading)
        if seg:
            segments.append(seg)
    return segments


def parse_markdown(path: Path, title: str | None = None) -> list[ParsedSegment]:
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise ParseError(f"could not read {path.name}: {exc}") from exc
    return parse_markdown_text(text, title or _title_from_path(path))


# -------------------------------------------------------------------- Transcripts


def parse_transcript_text(text: str, title: str) -> list[ParsedSegment]:
    """Parse a `.vtt`/`.srt` cue list or a timestamped plain-text transcript.

    Cues are grouped up to `_CUE_GROUP_CHARS` and labelled with the first
    timestamp in the group: "<title> @ 02:30".
    """
    lines = text.replace("\r\n", "\n").split("\n")
    segments: list[ParsedSegment] = []
    buffer: list[str] = []
    start: str | None = None

    def flush() -> None:
        nonlocal start
        ref = f"{title} @ {start}" if start else title
        seg = _segment(" ".join(buffer), ref, locator=start)
        if seg:
            segments.append(seg)
        buffer.clear()
        start = None

    for raw in lines:
        line = raw.strip()
        if not line or line.upper().startswith("WEBVTT") or line.startswith("NOTE "):
            continue
        timing = _CUE_TIMING.match(line)
        if timing:
            if start is not None and sum(len(part) for part in buffer) >= _CUE_GROUP_CHARS:
                flush()
            if start is None:
                start = _normalize_timestamp(timing.group("start"))
            continue
        if line.isdigit() and not buffer:
            continue  # SRT cue number
        inline = _INLINE_TIMESTAMP.match(line)
        if inline:
            if start is not None and sum(len(part) for part in buffer) >= _CUE_GROUP_CHARS:
                flush()
            if start is None:
                start = _normalize_timestamp(inline.group("ts"))
            line = line[inline.end() :].strip()
            if not line:
                continue
        buffer.append(line)
    flush()
    return segments


def parse_transcript(path: Path, title: str | None = None) -> list[ParsedSegment]:
    """Transcript entry point. Markdown transcripts route to the markdown parser.

    The seed pack ships its video transcript as markdown with authored chunk
    metadata, which already carries "… Video @ 02:30" refs.
    """
    doc_title = title or _title_from_path(path)
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise ParseError(f"could not read {path.name}: {exc}") from exc

    # A .md/.txt transcript with no cue timings and no inline timestamps is just
    # a document; authored chunk metadata always wins over timestamp sniffing.
    plain_document = (
        path.suffix.lower() in TEXT_SUFFIXES
        and not _CUE_TIMING.search(text)
        and (_CHUNK_COMMENT.search(text) or not _INLINE_TIMESTAMP.search(text))
    )
    if plain_document:
        return parse_markdown_text(text, doc_title)
    return parse_transcript_text(text, doc_title)


# --------------------------------------------------------------------------- URL


def check_fetch_target(url: str) -> None:
    """Refuse a URL pointing at this machine or the private network (SSRF).

    What this fetches is indexed and later quoted to a student, so a URL is not
    just a request — it is content injection performed from the server's own
    network position. Unchecked, `http://169.254.169.254/…` or an intranet host
    becomes course material.

    Hostnames are resolved and the *addresses* checked, because a public name
    can point anywhere; `parse_url` re-checks every redirect hop for the same
    reason.
    """
    parsed = urlparse(url)
    if parsed.scheme not in ALLOWED_URL_SCHEMES:
        raise ParseError(
            f"refusing to fetch {url!r}: only {', '.join(sorted(ALLOWED_URL_SCHEMES))} "
            "URLs can be ingested."
        )
    host = parsed.hostname
    if not host:
        raise ParseError(f"refusing to fetch {url!r}: the URL has no host.")

    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror as exc:
        raise ParseError(f"could not resolve {host!r}: {exc}") from exc

    for info in infos:
        address = ipaddress.ip_address(info[4][0])
        if (
            address.is_private
            or address.is_loopback
            or address.is_link_local
            or address.is_reserved
            or address.is_multicast
            or address.is_unspecified
        ):
            raise ParseError(
                f"refusing to fetch {url!r}: {host} resolves to {address}, which is on this "
                "machine or a private network. Only public web pages can be ingested."
            )


def parse_url(
    url: str, timeout: float = 20.0, client: httpx.Client | None = None
) -> list[ParsedSegment]:
    """Fetch a page and extract its main content; `source_ref` = "<page title> — <url>"."""
    from bs4 import BeautifulSoup

    check_fetch_target(url)

    owned = client is None
    # Redirects are followed by hand so each hop is re-checked: a public URL that
    # redirects to 127.0.0.1 would otherwise walk straight past the first check.
    http = client or httpx.Client(timeout=timeout, follow_redirects=False)
    try:
        current = url
        for _ in range(MAX_REDIRECTS):
            response = http.get(current, headers={"User-Agent": "socratic-tutor-poc/0.1"})
            if not response.is_redirect or not response.next_request:
                break
            current = str(response.next_request.url)
            check_fetch_target(current)
        else:
            raise ParseError(f"could not fetch {url}: more than {MAX_REDIRECTS} redirects.")

        response.raise_for_status()
        if len(response.content) > MAX_FETCH_BYTES:
            raise ParseError(
                f"refusing {url}: the page exceeds {MAX_FETCH_BYTES // (1024 * 1024)} MB."
            )
        html = response.text
    except httpx.HTTPError as exc:
        raise ParseError(f"could not fetch {url}: {exc}") from exc
    finally:
        if owned:
            http.close()

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
        tag.decompose()

    page_title = (soup.title.get_text(strip=True) if soup.title else "") or url
    root = soup.find("main") or soup.find("article") or soup.body or soup
    ref = f"{page_title} — {url}"

    segments: list[ParsedSegment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        seg = _segment(
            "\n\n".join(buffer),
            f"{ref} · {heading}" if heading else ref,
            locator=heading or url,
        )
        if seg:
            segments.append(seg)
        buffer.clear()

    for element in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote"]):
        text = element.get_text(" ", strip=True)
        if not text:
            continue
        if element.name in {"h1", "h2", "h3", "h4"}:
            flush()
            heading = text
            continue
        buffer.append(text)
    flush()

    if not segments:
        seg = _segment(root.get_text("\n", strip=True), ref, locator=url)
        if seg:
            segments.append(seg)
    return segments


# ---------------------------------------------------------------------- dispatch


def parse_path(path: str | Path, title: str | None = None) -> list[ParsedSegment]:
    """Parse any supported file by extension.

    Raises `ParseError` for a missing file or an unsupported format; `ingest`
    turns that into a per-file error rather than failing a batch.
    """
    path = Path(path)
    if not path.is_file():
        raise ParseError(f"not a file: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path, title)
    if suffix == ".docx":
        return parse_docx(path, title)
    if suffix == ".pptx":
        return parse_pptx(path, title)
    if suffix in TRANSCRIPT_SUFFIXES:
        return parse_transcript(path, title)
    raise ParseError(
        f"unsupported format '{suffix or path.name}' — supported: "
        f"{', '.join(sorted(SUPPORTED_SUFFIXES))}"
    )


__all__ = [
    "ALLOWED_URL_SCHEMES",
    "MAX_FETCH_BYTES",
    "MAX_REDIRECTS",
    "SUPPORTED_SUFFIXES",
    "ParseError",
    "ParsedSegment",
    "check_fetch_target",
    "detect_lang",
    "parse_docx",
    "parse_markdown",
    "parse_markdown_text",
    "parse_path",
    "parse_pdf",
    "parse_pptx",
    "parse_transcript",
    "parse_transcript_text",
    "parse_url",
]
